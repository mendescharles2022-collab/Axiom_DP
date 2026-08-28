"""
Gera o .docx do recibo avulso (contracheque/pró-labore) a partir do
resultado de app.services.calculo_folha.montar_calculo_recibo().

Ao contrário dos 48 modelos de documento (docxtpl, catálogo em
app/docs_templates/), este recibo tem uma estrutura própria — tabela de
rubricas de tamanho variável, bases de cálculo, duas vias (empregador e
empregado) — então é montado diretamente com python-docx em vez de um
template fixo. A estrutura de campos segue a planilha de referência
fornecida pelo escritório
(dados_para_importar/Recibo_de_Salário_-_Sebastião_das_Dores_da_Silva.xls
— HANDOFF seção 5.6).
"""
import os
from datetime import datetime

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

from app.config import OUTPUT_DIR
from app.utils.texto import titulo_pt

MESES_PT = ["janeiro", "fevereiro", "março", "abril", "maio", "junho", "julho",
            "agosto", "setembro", "outubro", "novembro", "dezembro"]

FRASE_PADRAO = "Declaro que recebi o valor descrito e dou plena quitação do que me é devido até o momento."


class GeracaoReciboError(Exception):
    pass


def _fmt(valor) -> str:
    if valor is None:
        return "-"
    return f"{valor:,.2f}".replace(",", "_").replace(".", ",").replace("_", ".")


def _documento_empregador(empresa):
    if empresa.cno:
        return "CNO", empresa.cno
    if empresa.cei:
        return "CEI", empresa.cei
    if empresa.caepf:
        return "CAEPF", empresa.caepf
    return empresa.tipo_inscricao or "CNPJ", empresa.cnpj


def _titulo_documento(recibo):
    return "Recibo de Pagamento de Pró-Labore" if recibo.tipo == "pro_labore" else "Recibo de Pagamento de Salário"


def _nome_beneficiario(recibo):
    if recibo.empregado:
        return recibo.empregado.nome_completo
    return recibo.nome_pro_labore or "-"


def _competencia_extenso(competencia: str) -> str:
    ano, mes = competencia.split("-")
    return f"{MESES_PT[int(mes) - 1]}/{ano}"


def _escrever_via(doc: Document, recibo, resultado, frase: str, rotulo_via: str):
    empresa = recibo.empresa
    doc_tipo, doc_numero = _documento_empregador(empresa)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = p.add_run(frase)
    run.italic = True
    run.font.size = Pt(8)

    titulo = doc.add_heading(_titulo_documento(recibo), level=2)

    tabela_cab = doc.add_table(rows=3, cols=2)
    tabela_cab.style = "Table Grid"
    tabela_cab.cell(0, 0).text = "Empregador"
    tabela_cab.cell(0, 1).text = titulo_pt(empresa.razao_social)
    tabela_cab.cell(1, 0).text = "Endereço"
    tabela_cab.cell(1, 1).text = empresa.endereco_completo()
    tabela_cab.cell(2, 0).text = doc_tipo
    tabela_cab.cell(2, 1).text = doc_numero or "-"

    doc.add_paragraph(f"Referente ao mês/ano: {_competencia_extenso(recibo.competencia)}")

    tabela_func = doc.add_table(rows=2, cols=3)
    tabela_func.style = "Table Grid"
    hdr = tabela_func.rows[0].cells
    hdr[0].text, hdr[1].text, hdr[2].text = "Código", "Nome", "Cargo/Função"
    dados = tabela_func.rows[1].cells
    dados[0].text = str(recibo.empregado_id or "-")
    dados[1].text = titulo_pt(_nome_beneficiario(recibo))
    dados[2].text = (recibo.empregado.cargo if recibo.empregado and recibo.empregado.cargo else "-")

    doc.add_paragraph()
    tabela_itens = doc.add_table(rows=1, cols=5)
    tabela_itens.style = "Table Grid"
    hdr = tabela_itens.rows[0].cells
    for i, titulo_col in enumerate(["Cód.", "Descrição", "Referência", "Proventos", "Descontos"]):
        hdr[i].text = titulo_col
    for item in recibo.itens:
        linha = tabela_itens.add_row().cells
        linha[0].text = str(item.rubrica.codigo)
        linha[1].text = item.rubrica.nome
        linha[2].text = item.referencia or "-"
        linha[3].text = _fmt(item.valor_provento) if item.valor_provento else ""
        linha[4].text = _fmt(item.valor_desconto) if item.valor_desconto else ""

    doc.add_paragraph()
    tabela_totais = doc.add_table(rows=2, cols=3)
    tabela_totais.style = "Table Grid"
    tabela_totais.rows[0].cells[0].text = "Total dos Vencimentos"
    tabela_totais.rows[0].cells[1].text = "Total dos Descontos"
    tabela_totais.rows[0].cells[2].text = "Líquido a Receber"
    tabela_totais.rows[1].cells[0].text = _fmt(resultado["total_proventos"])
    tabela_totais.rows[1].cells[1].text = _fmt(resultado["total_descontos"] + resultado["inss"]["valor"] + resultado["irrf"]["valor_final"])
    tabela_totais.rows[1].cells[2].text = _fmt(resultado["liquido"])

    doc.add_paragraph()
    tabela_bases = doc.add_table(rows=2, cols=6)
    tabela_bases.style = "Table Grid"
    rotulos_base = ["Base Cálc. INSS", "INSS", "Base Cálc. FGTS", "FGTS do Mês", "Base Cálc. IRRF", "IRRF"]
    valores_base = [
        _fmt(resultado["base_inss"]), _fmt(resultado["inss"]["valor"]),
        _fmt(resultado["base_fgts"]), _fmt(resultado["fgts"]),
        _fmt(resultado["base_irrf"]), _fmt(resultado["irrf"]["valor_final"]),
    ]
    for i, rotulo in enumerate(rotulos_base):
        tabela_bases.rows[0].cells[i].text = rotulo
        tabela_bases.rows[1].cells[i].text = valores_base[i]

    doc.add_paragraph()
    p = doc.add_paragraph()
    p.add_run("Data: ____/____/________" + " " * 20 + "Assinatura: _______________________________")

    p = doc.add_paragraph(rotulo_via)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in p.runs:
        run.bold = True
        run.font.size = Pt(9)


def _slug(texto: str) -> str:
    import re
    import unicodedata
    texto = unicodedata.normalize("NFKD", texto or "").encode("ascii", "ignore").decode()
    texto = re.sub(r"[^A-Za-z0-9]+", "_", texto).strip("_")
    return texto or "recibo"


def gerar_recibo_docx(recibo, resultado) -> str:
    frase = recibo.frase_quitacao.texto if recibo.frase_quitacao else FRASE_PADRAO

    doc = Document()
    _escrever_via(doc, recibo, resultado, frase, "1ª VIA - EMPREGADOR")
    doc.add_page_break()
    _escrever_via(doc, recibo, resultado, frase, "2ª VIA - EMPREGADO")

    os.makedirs(OUTPUT_DIR, exist_ok=True)
    nome_beneficiario = _slug(_nome_beneficiario(recibo))
    carimbo = datetime.now().strftime("%Y%m%d_%H%M%S")
    nome_arquivo = f"recibo_{recibo.tipo}_{nome_beneficiario}_{recibo.competencia}_{carimbo}.docx"
    caminho = os.path.join(OUTPUT_DIR, nome_arquivo)
    doc.save(caminho)
    return caminho

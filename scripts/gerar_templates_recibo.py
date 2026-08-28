"""
Gera os modelos visuais de recibo avulso (contracheque e pró-labore) em
app/docs_templates_recibo/ — script de autoria, roda uma vez (ou de novo
se algum modelo precisar ser refeito do zero). Segue a mesma identidade
visual dos 48 modelos de DP já existentes (Cambria navy #1F3864 para
título/nome da empresa, Calibri cinza #595959 para metadados, corpo em
Calibri) — ver HANDOFF ADENDO seção A.

Cada modelo é um .docx com placeholders docxtpl ({{ }}) e um loop de
linha de tabela ({%tr for item in itens %} ... {%tr endfor %}) para a
lista de rubricas, que app/services/recibo_engine.py preenche.

Uso: python scripts/gerar_templates_recibo.py
"""
import os

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

RAIZ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DESTINO = os.path.join(RAIZ, "app", "docs_templates_recibo")

NAVY = RGBColor(0x1F, 0x38, 0x64)
NAVY_ESCURO = RGBColor(0x0F, 0x1E, 0x3D)
DOURADO = RGBColor(0xB0, 0x8D, 0x3E)
CINZA = RGBColor(0x59, 0x59, 0x59)
BRANCO = RGBColor(0xFF, 0xFF, 0xFF)


def _sombrear(cell, cor_hex: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), cor_hex)
    tc_pr.append(shd)


def _run(paragrafo, texto, *, negrito=False, cor=None, tamanho=11, fonte="Calibri", italico=False):
    r = paragrafo.add_run(texto)
    r.bold = negrito
    r.italic = italico
    r.font.size = Pt(tamanho)
    r.font.name = fonte
    if cor:
        r.font.color.rgb = cor
    return r


def _celula(cell, texto, *, negrito=False, cor=None, tamanho=10, alinhar_direita=False, fonte="Calibri"):
    cell.paragraphs[0].text = ""
    if alinhar_direita:
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
    _run(cell.paragraphs[0], texto, negrito=negrito, cor=cor, tamanho=tamanho, fonte=fonte)


def _linha_controle(tabela, tag: str, ncols: int):
    """
    Linha de controle para loop/condicional de linha no docxtpl: a tag
    ({%tr for ...%}, {%tr endfor%}, {%tr if ...%}, {%tr endif%}) precisa
    estar SOZINHA numa linha própria — docxtpl colapsa a linha inteira
    que contém a tag "trX" no tag puro ({% for %}/{% endfor %}/...),
    então o conteúdo real (os dados) tem que ficar numa linha SEPARADA
    entre a linha do "for"/"if" e a do "endfor"/"endif".
    """
    linha = tabela.add_row().cells
    _celula(linha[0], tag)
    for i in range(1, ncols):
        _celula(linha[i], "")


def _largura_pagina_estreita(doc):
    for section in doc.sections:
        section.left_margin = Cm(1.8)
        section.right_margin = Cm(1.8)
        section.top_margin = Cm(1.5)
        section.bottom_margin = Cm(1.5)


def _rodape_padrao(doc):
    for section in doc.sections:
        p = section.footer.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _run(p, "Documento de uso interno do Departamento Pessoal.", cor=CINZA, tamanho=8)


# ------------------------------------------------------------- Clássico

def _cabecalho_classico(doc, titulo_documento):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    _run(p, "{{ frase_quitacao }}", italico=True, tamanho=8, cor=CINZA)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _run(p, "{{ empresa.razao_social }}", negrito=True, cor=NAVY, tamanho=14, fonte="Cambria")

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _run(p, "{{ empresa.documento_rotulo }} {{ empresa.documento_numero }}  ·  {{ empresa.endereco_completo }}", cor=CINZA, tamanho=8)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _run(p, titulo_documento, negrito=True, cor=NAVY, tamanho=15, fonte="Cambria")

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _run(p, "Referente à competência de {{ competencia_extenso }}", cor=CINZA, tamanho=10)


def _tabela_beneficiario_classico(doc):
    tabela = doc.add_table(rows=2, cols=3)
    tabela.style = "Table Grid"
    hdr = tabela.rows[0].cells
    for i, t in enumerate(["Código", "Nome", "Cargo/Função"]):
        _celula(hdr[i], t, negrito=True)
        _sombrear(hdr[i], "EEF2FA")
    dados = tabela.rows[1].cells
    _celula(dados[0], "{{ beneficiario.codigo }}")
    _celula(dados[1], "{{ beneficiario.nome_completo }}")
    _celula(dados[2], "{{ beneficiario.cargo }}")


def _tabela_itens_classico(doc):
    doc.add_paragraph()
    tabela = doc.add_table(rows=1, cols=5)
    tabela.style = "Table Grid"
    hdr = tabela.rows[0].cells
    for i, t in enumerate(["Cód.", "Descrição", "Referência", "Proventos (R$)", "Descontos (R$)"]):
        _celula(hdr[i], t, negrito=True)
        _sombrear(hdr[i], "1F3864")
        hdr[i].paragraphs[0].runs[0].font.color.rgb = BRANCO

    _linha_controle(tabela, "{%tr for item in itens %}", 5)
    linha_loop = tabela.add_row().cells
    _celula(linha_loop[0], "{{ item.codigo }}")
    _celula(linha_loop[1], "{{ item.nome }}")
    _celula(linha_loop[2], "{{ item.referencia }}")
    _celula(linha_loop[3], "{{ item.provento }}")
    _celula(linha_loop[4], "{{ item.desconto }}")
    _linha_controle(tabela, "{%tr endfor %}", 5)

    _linha_controle(tabela, "{%tr if salario_familia.elegivel %}", 5)
    linha_sf = tabela.add_row().cells
    _celula(linha_sf[0], "-")
    _celula(linha_sf[1], "SALÁRIO-FAMÍLIA ({{ salario_familia.dependentes }} dependente(s) x {{ salario_familia.valor_cota }})")
    _celula(linha_sf[2], "-")
    _celula(linha_sf[3], "{{ salario_familia.valor }}")
    _celula(linha_sf[4], "")
    _linha_controle(tabela, "{%tr endif %}", 5)


def _tabela_totais_classico(doc):
    doc.add_paragraph()
    tabela = doc.add_table(rows=2, cols=3)
    tabela.style = "Table Grid"
    hdr = tabela.rows[0].cells
    for i, t in enumerate(["Total dos Vencimentos", "Total dos Descontos", "Líquido a Receber"]):
        _celula(hdr[i], t, negrito=True)
        _sombrear(hdr[i], "EEF2FA")
    dados = tabela.rows[1].cells
    _celula(dados[0], "{{ total_vencimentos }}", negrito=True)
    _celula(dados[1], "{{ total_descontos_com_tributos }}", negrito=True)
    _celula(dados[2], "{{ liquido }}", negrito=True, cor=NAVY)


def _tabela_bases_classico(doc, incluir_fgts):
    doc.add_paragraph()
    rotulos = ["Base Cálc. INSS", "INSS", "Base Cálc. IRRF", "IRRF (faixa {{ irrf_aliquota }}%)"]
    valores = ["{{ inss_base }}", "{{ inss_valor }}", "{{ irrf_base }}", "{{ irrf_valor }}"]
    if incluir_fgts:
        rotulos += ["Base Cálc. FGTS", "FGTS do Mês"]
        valores += ["{{ fgts_base }}", "{{ fgts_valor }}"]
    tabela = doc.add_table(rows=2, cols=len(rotulos))
    tabela.style = "Table Grid"
    for i, r in enumerate(rotulos):
        _celula(tabela.rows[0].cells[i], r, negrito=True, tamanho=9)
        _sombrear(tabela.rows[0].cells[i], "EEF2FA")
        _celula(tabela.rows[1].cells[i], valores[i], tamanho=9)


def _assinatura_classico(doc):
    doc.add_paragraph()
    p = doc.add_paragraph()
    _run(p, "Data: ____/____/________" + " " * 16 + "Assinatura: _______________________________________")


def _construir_classico(tipo: str, titulo_documento: str, incluir_fgts: bool, rotulo_via: str | None = None):
    doc = Document()
    _largura_pagina_estreita(doc)
    _rodape_padrao(doc)

    _cabecalho_classico(doc, titulo_documento)
    _tabela_beneficiario_classico(doc)
    _tabela_itens_classico(doc)
    _tabela_totais_classico(doc)
    _tabela_bases_classico(doc, incluir_fgts)
    _assinatura_classico(doc)

    if rotulo_via:
        p = doc.add_paragraph(rotulo_via)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.runs[0].bold = True
        p.runs[0].font.size = Pt(9)
        p.runs[0].font.color.rgb = CINZA

        doc.add_page_break()
        _cabecalho_classico(doc, titulo_documento)
        _tabela_beneficiario_classico(doc)
        _tabela_itens_classico(doc)
        _tabela_totais_classico(doc)
        _tabela_bases_classico(doc, incluir_fgts)
        _assinatura_classico(doc)
        p = doc.add_paragraph(rotulo_via.replace("1ª", "2ª").replace("EMPREGADOR", "EMPREGADO"))
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.runs[0].bold = True
        p.runs[0].font.size = Pt(9)
        p.runs[0].font.color.rgb = CINZA

    return doc


# -------------------------------------------------------------- Moderno

def _faixa_titulo_moderno(doc, titulo_documento):
    tabela = doc.add_table(rows=1, cols=1)
    tabela.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tabela.rows[0].cells[0]
    _sombrear(cell, "1F3864")
    cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    _run(cell.paragraphs[0], "{{ empresa.razao_social }}", negrito=True, cor=BRANCO, tamanho=15, fonte="Cambria")
    p2 = cell.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _run(p2, "{{ empresa.documento_rotulo }} {{ empresa.documento_numero }}  ·  {{ empresa.endereco_completo }}", cor=RGBColor(0xC9, 0xD3, 0xE6), tamanho=8)
    p3 = cell.add_paragraph()
    p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _run(p3, titulo_documento.upper(), negrito=True, cor=RGBColor(0xE8, 0xC7, 0x7A), tamanho=11)

    doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _run(p, "Competência {{ competencia_extenso }}", cor=CINZA, tamanho=10, italico=True)


def _linha_beneficiario_moderno(doc):
    p = doc.add_paragraph()
    _run(p, "Beneficiário: ", negrito=True, tamanho=10)
    _run(p, "{{ beneficiario.nome_completo }}", tamanho=10)
    _run(p, "   ·   Código: ", negrito=True, tamanho=10)
    _run(p, "{{ beneficiario.codigo }}", tamanho=10)
    _run(p, "   ·   Cargo: ", negrito=True, tamanho=10)
    _run(p, "{{ beneficiario.cargo }}", tamanho=10)


def _tabela_itens_moderno(doc):
    doc.add_paragraph()
    tabela = doc.add_table(rows=1, cols=5)
    tabela.style = "Table Grid"
    hdr = tabela.rows[0].cells
    for i, t in enumerate(["Cód.", "Descrição", "Referência", "Proventos (R$)", "Descontos (R$)"]):
        _celula(hdr[i], t, negrito=True, cor=BRANCO)
        _sombrear(hdr[i], "B08D3E")

    _linha_controle(tabela, "{%tr for item in itens %}", 5)
    linha_loop = tabela.add_row().cells
    _celula(linha_loop[0], "{{ item.codigo }}")
    _celula(linha_loop[1], "{{ item.nome }}")
    _celula(linha_loop[2], "{{ item.referencia }}")
    _celula(linha_loop[3], "{{ item.provento }}")
    _celula(linha_loop[4], "{{ item.desconto }}")
    _linha_controle(tabela, "{%tr endfor %}", 5)

    _linha_controle(tabela, "{%tr if salario_familia.elegivel %}", 5)
    linha_sf = tabela.add_row().cells
    _celula(linha_sf[0], "-")
    _celula(linha_sf[1], "SALÁRIO-FAMÍLIA ({{ salario_familia.dependentes }} dependente(s) x {{ salario_familia.valor_cota }})")
    _celula(linha_sf[2], "-")
    _celula(linha_sf[3], "{{ salario_familia.valor }}")
    _celula(linha_sf[4], "")
    _linha_controle(tabela, "{%tr endif %}", 5)


def _resumo_moderno(doc, incluir_fgts):
    doc.add_paragraph()
    tabela = doc.add_table(rows=2, cols=3)
    tabela.style = "Table Grid"
    hdr = tabela.rows[0].cells
    for i, t in enumerate(["Total dos Vencimentos", "Total dos Descontos", "Líquido a Receber"]):
        _celula(hdr[i], t, negrito=True, cor=BRANCO)
        _sombrear(hdr[i], "1F3864")
    dados = tabela.rows[1].cells
    _celula(dados[0], "{{ total_vencimentos }}", negrito=True)
    _celula(dados[1], "{{ total_descontos_com_tributos }}", negrito=True)
    _celula(dados[2], "{{ liquido }}", negrito=True, cor=NAVY, tamanho=12)

    doc.add_paragraph()
    p = doc.add_paragraph()
    _run(p, "Base INSS ", negrito=True, tamanho=9)
    _run(p, "R$ {{ inss_base }}", tamanho=9)
    _run(p, "  →  INSS ", negrito=True, tamanho=9)
    _run(p, "R$ {{ inss_valor }}", tamanho=9)
    _run(p, "   |   Base IRRF ", negrito=True, tamanho=9)
    _run(p, "R$ {{ irrf_base }}", tamanho=9)
    _run(p, "  →  IRRF (faixa {{ irrf_aliquota }}%) ", negrito=True, tamanho=9)
    _run(p, "R$ {{ irrf_valor }}", tamanho=9)
    if incluir_fgts:
        _run(p, "   |   Base FGTS ", negrito=True, tamanho=9)
        _run(p, "R$ {{ fgts_base }}", tamanho=9)
        _run(p, "  →  FGTS ", negrito=True, tamanho=9)
        _run(p, "R$ {{ fgts_valor }}", tamanho=9)


def _rodape_quitacao_moderno(doc):
    doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _run(p, "{{ frase_quitacao }}", italico=True, tamanho=9, cor=CINZA)

    doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _run(p, "Data: ____/____/________" + " " * 16 + "Assinatura: _______________________________________", tamanho=10)


def _construir_moderno(titulo_documento: str, incluir_fgts: bool):
    doc = Document()
    _largura_pagina_estreita(doc)
    _rodape_padrao(doc)

    _faixa_titulo_moderno(doc, titulo_documento)
    _linha_beneficiario_moderno(doc)
    _tabela_itens_moderno(doc)
    _resumo_moderno(doc, incluir_fgts)
    _rodape_quitacao_moderno(doc)
    return doc


def main():
    os.makedirs(DESTINO, exist_ok=True)

    modelos = [
        ("contracheque_classico.docx",
         _construir_classico("contracheque", "Recibo de Pagamento de Salário", True, "1ª VIA - EMPREGADOR")),
        ("contracheque_moderno.docx",
         _construir_moderno("Recibo de Pagamento de Salário", True)),
        ("pro_labore_classico.docx",
         _construir_classico("pro_labore", "Recibo de Pagamento de Pró-Labore", False, "1ª VIA - EMPRESA")),
        ("pro_labore_moderno.docx",
         _construir_moderno("Recibo de Pagamento de Pró-Labore", False)),
    ]

    for nome_arquivo, doc in modelos:
        caminho = os.path.join(DESTINO, nome_arquivo)
        doc.save(caminho)
        print(f"Gerado: {caminho}")


if __name__ == "__main__":
    main()

import os

from docx import Document

from app.extensions import db
from app.models.empresa import Empresa
from app.models.empregado import Empregado
from app.models.frase_quitacao import FraseQuitacao
from app.models.recibo_avulso import ReciboAvulso
from app.models.rubrica import Rubrica
from app.services.calculo_folha import montar_calculo_recibo
from app.services.recibo_engine import gerar_recibo_docx


def _seed_tabelas(app):
    with app.app_context():
        from scripts.seed_tabelas_fiscais import seed_inss, seed_irrf, seed_irrf_redutor

        seed_inss()
        seed_irrf()
        seed_irrf_redutor()


def _criar_empresa_empregado_rubricas(app):
    with app.app_context():
        empresa = Empresa(razao_social="Empresa Recibo LTDA", cnpj="66.666.666/0001-66", cei="38.730.06578/89")
        db.session.add(empresa)
        db.session.commit()
        empregado = Empregado(
            empresa_id=empresa.id, nome_completo="TRABALHADOR DE TESTE", cpf="111.222.333-44",
            cargo="Auxiliar",
        )
        db.session.add(empregado)

        salario = Rubrica(codigo=8781, nome="SALARIO EMPREGADO", tipo="P",
                           incidencia_irrf=11, incidencia_inss=11, incidencia_fgts=11, incidencia_pis=11)
        vt = Rubrica(codigo=999, nome="VALE TRANSPORTE", tipo="D",
                     incidencia_irrf=0, incidencia_inss=0, incidencia_fgts=0, incidencia_pis=0)
        db.session.add_all([salario, vt])
        db.session.commit()
        return empresa.id, empregado.id, salario.id, vt.id, salario.codigo, vt.codigo


def test_criar_recibo_via_formulario_e_calcular(app, auth_client):
    _seed_tabelas(app)
    empresa_id, empregado_id, salario_id, vt_id, cod_salario, cod_vt = _criar_empresa_empregado_rubricas(app)

    dados = {
        "tipo": "contracheque",
        "competencia": "2024-06",
        "empregado_id": str(empregado_id),
        "rubrica_texto_0": f"{cod_salario} — SALARIO EMPREGADO",
        "valor_provento_0": "3.000,00",
        "referencia_0": "30",
        "rubrica_texto_1": str(cod_vt),
        "valor_desconto_1": "180,00",
    }
    resp = auth_client.post(f"/empresas/{empresa_id}/recibos/novo", data=dados, follow_redirects=True)
    assert resp.status_code == 200
    corpo = resp.get_data(as_text=True)
    assert "258.82" in corpo  # INSS calculado sobre R$ 3.000,00 (ver test_calculo_folha)

    with app.app_context():
        recibo = ReciboAvulso.query.filter_by(empresa_id=empresa_id).first()
        assert recibo is not None
        assert len(recibo.itens) == 2


def test_recibo_sem_itens_e_rejeitado(app, auth_client):
    _seed_tabelas(app)
    empresa_id, empregado_id, *_resto = _criar_empresa_empregado_rubricas(app)

    resp = auth_client.post(
        f"/empresas/{empresa_id}/recibos/novo",
        data={"tipo": "contracheque", "competencia": "2024-06", "empregado_id": str(empregado_id)},
        follow_redirects=True,
    )
    assert resp.status_code == 200
    assert "Adicione pelo menos um item" in resp.get_data(as_text=True)

    with app.app_context():
        assert ReciboAvulso.query.count() == 0


def test_pro_labore_sem_empregado_usa_nome_informado(app, auth_client):
    _seed_tabelas(app)
    empresa_id, _empregado_id, salario_id, _vt_id, cod_salario, _cod_vt = _criar_empresa_empregado_rubricas(app)

    resp = auth_client.post(
        f"/empresas/{empresa_id}/recibos/novo",
        data={
            "tipo": "pro_labore", "competencia": "2024-06", "nome_pro_labore": "SOCIO EXEMPLO",
            "rubrica_texto_0": str(cod_salario), "valor_provento_0": "2.000,00",
        },
        follow_redirects=True,
    )
    assert resp.status_code == 200

    with app.app_context():
        recibo = ReciboAvulso.query.filter_by(tipo="pro_labore").first()
        assert recibo.empregado_id is None
        assert recibo.nome_pro_labore == "SOCIO EXEMPLO"


def test_gerar_e_baixar_documento_do_recibo(app, auth_client):
    _seed_tabelas(app)
    empresa_id, empregado_id, salario_id, vt_id, cod_salario, cod_vt = _criar_empresa_empregado_rubricas(app)

    auth_client.post(
        f"/empresas/{empresa_id}/recibos/novo",
        data={
            "tipo": "contracheque", "competencia": "2024-06", "empregado_id": str(empregado_id),
            "rubrica_texto_0": str(cod_salario), "valor_provento_0": "3000",
            "rubrica_texto_1": str(cod_vt), "valor_desconto_1": "180",
        },
    )
    with app.app_context():
        recibo_id = ReciboAvulso.query.first().id

    resp = auth_client.post(f"/recibos/{recibo_id}/gerar", follow_redirects=True)
    assert resp.status_code == 200
    assert "gerado com sucesso" in resp.get_data(as_text=True)

    resp_baixar = auth_client.get(f"/recibos/{recibo_id}/baixar")
    assert resp_baixar.status_code == 200
    assert resp_baixar.headers["Content-Type"].startswith(
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    )


def test_gerar_recibo_docx_contem_dados_esperados(app):
    _seed_tabelas(app)
    empresa_id, empregado_id, salario_id, vt_id, _cod_salario, _cod_vt = _criar_empresa_empregado_rubricas(app)

    with app.app_context():
        from app.models.recibo_avulso import ReciboAvulsoItem

        recibo = ReciboAvulso(empresa_id=empresa_id, empregado_id=empregado_id, competencia="2024-06", tipo="contracheque")
        db.session.add(recibo)
        db.session.commit()
        db.session.add_all([
            ReciboAvulsoItem(recibo_id=recibo.id, rubrica_id=salario_id, valor_provento="3000.00", referencia="30"),
            ReciboAvulsoItem(recibo_id=recibo.id, rubrica_id=vt_id, valor_desconto="180.00"),
        ])
        db.session.commit()

        resultado = montar_calculo_recibo(recibo)
        caminho = gerar_recibo_docx(recibo, resultado)

        doc = Document(caminho)
        todo_texto = "\n".join(p.text for p in doc.paragraphs)
        for t in doc.tables:
            for row in t.rows:
                for cell in row.cells:
                    todo_texto += "\n" + cell.text

    assert "Trabalhador de Teste" in todo_texto
    assert "1ª VIA - EMPREGADOR" in todo_texto
    assert "2ª VIA - EMPREGADO" in todo_texto
    assert "258,82" in todo_texto  # INSS (formatado em pt-BR pelo recibo_engine._fmt)
    assert os.path.exists(caminho)


def test_frase_quitacao_padrao_da_empresa_e_usada(app, auth_client):
    _seed_tabelas(app)
    empresa_id, empregado_id, salario_id, _vt_id, cod_salario, _cod_vt = _criar_empresa_empregado_rubricas(app)

    with app.app_context():
        frase = FraseQuitacao(nome="Cliente Especial", texto="Frase customizada de teste.")
        db.session.add(frase)
        db.session.commit()
        empresa = Empresa.query.get(empresa_id)
        empresa.frase_quitacao_padrao_id = frase.id
        db.session.commit()

    auth_client.post(
        f"/empresas/{empresa_id}/recibos/novo",
        data={
            "tipo": "contracheque", "competencia": "2024-06", "empregado_id": str(empregado_id),
            "rubrica_texto_0": str(cod_salario), "valor_provento_0": "1000",
        },
    )
    with app.app_context():
        recibo = ReciboAvulso.query.filter_by(empresa_id=empresa_id).first()
        assert recibo.frase_quitacao.texto == "Frase customizada de teste."

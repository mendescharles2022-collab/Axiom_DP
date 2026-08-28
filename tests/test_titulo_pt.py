import pytest

from app.utils.texto import titulo_pt


@pytest.mark.parametrize(
    "bruto,esperado",
    [
        (None, None),
        ("", ""),
        ("EMPRESA EXEMPLO DE COMERCIO LTDA", "Empresa Exemplo de Comercio LTDA"),
        ("JOAO DA SILVA E SOUZA ME", "Joao da Silva e Souza ME"),
        ("CASA DAS FLORES EIRELI", "Casa das Flores EIRELI"),
        ("COMERCIAL DOS SANTOS S/A", "Comercial dos Santos S/A"),
        ("MARIA JOSE-CARLOS PEREIRA", "Maria Jose-Carlos Pereira"),
        ("EMPRESA XPTO MEI", "Empresa Xpto MEI"),
        ("de oliveira comercio", "De Oliveira Comercio"),  # conector nunca é a 1ª palavra
    ],
)
def test_titulo_pt(bruto, esperado):
    assert titulo_pt(bruto) == esperado


def test_endereco_completo_aplica_titulo_pt(app):
    from app.models.empresa import Empresa

    with app.app_context():
        empresa = Empresa(
            razao_social="EMPRESA TESTE LTDA",
            cnpj="11.111.111/0001-11",
            logradouro="RUA DAS ACACIAS",
            numero="100",
            bairro="CENTRO",
            municipio="ITAPACI",
            uf="GO",
            cep="76490-000",
        )
        endereco = empresa.endereco_completo()

    assert "Rua das Acacias" in endereco
    assert "Centro" in endereco
    assert "Itapaci/GO" in endereco
    assert "CEP 76490-000" in endereco  # "CEP" continua maiúsculo (rótulo fixo)


def test_documento_gerado_com_titulo_pt(app, auth_client):
    from app.extensions import db
    from app.models.empresa import Empresa
    from app.models.empregado import Empregado
    from app.models.documento import TemplateDocumento
    from tests.test_smoke import _seed_templates

    with app.app_context():
        _seed_templates()
        empresa = Empresa(razao_social="EMPRESA CAIXA ALTA LTDA", cnpj="44.444.444/0001-44")
        db.session.add(empresa)
        db.session.commit()
        empregado = Empregado(
            empresa_id=empresa.id, nome_completo="FULANO DA SILVA", cpf="555.555.555-55"
        )
        db.session.add(empregado)
        db.session.commit()
        template = TemplateDocumento.query.first()
        empresa_id, empregado_id, template_id = empresa.id, empregado.id, template.id

    resp = auth_client.post(
        f"/empresas/{empresa_id}/emitir/{template_id}",
        data={"empregado_id": str(empregado_id)},
        follow_redirects=True,
    )
    assert resp.status_code == 200
    assert "gerado com sucesso" in resp.get_data(as_text=True)

    with app.app_context():
        from app.models.emissao import DocumentoEmitido
        from docx import Document

        doc_registro = DocumentoEmitido.query.order_by(DocumentoEmitido.id.desc()).first()
        documento = Document(doc_registro.caminho_arquivo_gerado)
        trechos = [p.text for p in documento.paragraphs]
        for secao in documento.sections:
            trechos += [p.text for p in secao.header.paragraphs]
            trechos += [p.text for p in secao.footer.paragraphs]
        texto = "\n".join(trechos)

    assert "EMPRESA CAIXA ALTA LTDA" not in texto
    assert "Empresa Caixa Alta LTDA" in texto
